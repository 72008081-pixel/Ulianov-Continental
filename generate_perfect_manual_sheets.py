# -*- coding: utf-8 -*-
"""
generate_perfect_manual_sheets.py
Builds the upgraded, 100% dynamic manual sheets (Tabs 2 & 3) supporting up to 8 BARS
and up to 6 NODES (12 DOFs) with automatic support detection, dynamic DOF assignment,
color tracking from local 4x4 to expanded 12x12 matrices, blank slot handling,
regularized matrix inversion (no #NUM!), reactions, and member forces.
Cátedra: Ing. Ulianov Cuba Valencia - Análisis Estructural II
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import CellIsRule

# Colors and Fonts matching course aesthetic
f_hero = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
f_hero_sub = Font(name="Segoe UI", size=9, italic=True, color="CBD5E1")
f_sec_hdr = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
f_tbl_hdr = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
f_tbl_sub = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
f_regular = Font(name="Segoe UI", size=9)
f_bold = Font(name="Segoe UI", size=9, bold=True)
f_input = Font(name="Segoe UI", size=9, bold=True, color="1E3A8A")
f_note = Font(name="Segoe UI", size=8, italic=True, color="64748B")
f_prof_title = Font(name="Segoe UI", size=10, bold=True, underline="single", color="1E293B")
f_prof_title_no_ul = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
f_prof_note = Font(name="Segoe UI", size=8, bold=True, color="334155")
f_yellow_banner = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
f_matrix_sym = Font(name="Segoe UI", size=11, bold=True, color="1E293B")
f_pedag_hdr = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
f_pedag_txt = Font(name="Segoe UI", size=8, color="334155")

fill_hero = PatternFill("solid", fgColor="1E293B")
fill_sec = PatternFill("solid", fgColor="334155")
fill_sec_teal = PatternFill("solid", fgColor="0F766E")
fill_tbl_sub = PatternFill("solid", fgColor="F1F5F9")
fill_input = PatternFill("solid", fgColor="FEF3C7") # Warm soft yellow input
fill_peach = PatternFill("solid", fgColor="FEE2E2") # Peach fill for C and D vectors
fill_prof_green = PatternFill("solid", fgColor="DCFCE7") # Soft sage green for free DOFs & K11
fill_yellow_banner = PatternFill("solid", fgColor="FEF08A") # Yellow banner for formulas
fill_gray_reac = PatternFill("solid", fgColor="EDEDFA") # Soft lavender for reactions & K21
fill_karm_diag = PatternFill("solid", fgColor="D1FAE5") # Uniform pastel green diagonal
fill_card = PatternFill("solid", fgColor="F8FAFC")
fill_white = PatternFill("solid", fgColor="FFFFFF")
fill_hdr_pedag = PatternFill("solid", fgColor="E2E8F0")
fill_card_pedag = PatternFill("solid", fgColor="F8FAFC")

border_thin = Border(
    left=Side(style='thin', color="CBD5E1"),
    right=Side(style='thin', color="CBD5E1"),
    top=Side(style='thin', color="CBD5E1"),
    bottom=Side(style='thin', color="CBD5E1")
)
border_pedag = Border(
    left=Side(style='medium', color="94A3B8"),
    right=Side(style='thin', color="CBD5E1"),
    top=Side(style='thin', color="CBD5E1"),
    bottom=Side(style='thin', color="CBD5E1")
)
border_diag_med = Border(
    left=Side(style='thin', color="CBD5E1"),
    right=Side(style='thin', color="CBD5E1"),
    top=Side(style='medium', color="10B981"),
    bottom=Side(style='medium', color="10B981")
)

al_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
al_left = Alignment(horizontal="left", vertical="center")
al_right = Alignment(horizontal="right", vertical="center")

BAR_PASTEL_PALETTE = [
    {"bg": "E0F2FE", "name": "Azul Cielo"},
    {"bg": "DCFCE7", "name": "Menta Suave"},
    {"bg": "FEE2E2", "name": "Rosa Coral"},
    {"bg": "FEF3C7", "name": "Melocotón"},
    {"bg": "F3E8FF", "name": "Lavanda"},
    {"bg": "CCFBF1", "name": "Turquesa"},
    {"bg": "FEF9C3", "name": "Mantequilla"},
    {"bg": "EDE7F6", "name": "Lila Pastel"},
]

def apply_styling(ws, r1, r2, c1, c2, fill=None, font=None, align=None, border=border_thin, num_fmt=None):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(r, c)
            if fill: cell.fill = fill
            if font: cell.font = font
            if align: cell.alignment = align
            if border is not None: cell.border = border
            if num_fmt: cell.number_format = num_fmt

def write_pedag_card(ws, r_start, r_end, c_start, c_end, title, lines):
    if r_end < r_start:
        r_end = r_start + len(lines)
    ws.merge_cells(start_row=r_start, start_column=c_start, end_row=r_start, end_column=c_end)
    ws.cell(r_start, c_start, title)
    apply_styling(ws, r_start, r_start, c_start, c_end, fill=fill_hdr_pedag, font=f_pedag_hdr, align=al_left, border=border_pedag)
    ws.row_dimensions[r_start].height = 20

    for i, line_txt in enumerate(lines):
        curr_r = r_start + 1 + i
        if curr_r <= r_end:
            ws.merge_cells(start_row=curr_r, start_column=c_start, end_row=curr_r, end_column=c_end)
            ws.cell(curr_r, c_start, line_txt)
            apply_styling(ws, curr_r, curr_r, c_start, c_end, fill=fill_card_pedag, font=f_pedag_txt, align=al_left, border=border_pedag)
            ws.row_dimensions[curr_r].height = 17

    for curr_r in range(r_start + 1 + len(lines), r_end + 1):
        ws.merge_cells(start_row=curr_r, start_column=c_start, end_row=curr_r, end_column=c_end)
        apply_styling(ws, curr_r, curr_r, c_start, c_end, fill=fill_card_pedag, font=f_pedag_txt, align=al_left, border=border_pedag)
        ws.row_dimensions[curr_r].height = 17


def build_teacher_manual_sheet(wb, sheet_name, is_isostatic=True):
    """
    Builds the manual template sheet supporting up to 8 BARS and up to 6 NODES (12 DOFs)
    following EXACTLY the classroom layout of Ing. Ulianov Cuba Valencia.
    Fully automated boundary condition detection, color-tracked expanded matrices,
    regularized K11/K21, and zero-error inversion.
    """
    if sheet_name in wb.sheetnames:
        idx = wb.sheetnames.index(sheet_name)
        wb.remove(wb[sheet_name])
        ws = wb.create_sheet(title=sheet_name, index=idx)
    else:
        ws = wb.create_sheet(title=sheet_name)

    ws.views.sheetView[0].showGridLines = True

    MAX_BARS = 8
    MAX_NODES = 6
    dofs_total = 2 * MAX_NODES # 12 DOFs (U1X..U6Y)

    if is_isostatic:
        # Default: User's benchmark with 3 nodes active, nodes 4-6 ready for input
        nodes_data = [
            {"id": 1, "x": 0.0,    "y": 0.0,   "support": "Móvil Y", "px": 0.0,    "py": 0.0},
            {"id": 2, "x": -700.0, "y": 0.0,   "support": "Fijo",    "px": 0.0,    "py": 0.0},
            {"id": 3, "x": -400.0, "y": 500.0, "support": "Libre",   "px": 4000.0, "py": -5000.0},
            {"id": 4, "x": None,   "y": None,  "support": "Libre",   "px": None,   "py": None},
            {"id": 5, "x": None,   "y": None,  "support": "Libre",   "px": None,   "py": None},
            {"id": 6, "x": None,   "y": None,  "support": "Libre",   "px": None,   "py": None},
        ]
        bars_data = [
            {"id": 1, "start": 1, "end": 3, "a": 10.0, "e": 2100000.0},
            {"id": 2, "start": 2, "end": 3, "a": 10.0, "e": 2100000.0},
            {"id": 3, "start": 2, "end": 1, "a": 10.0, "e": 2100000.0},
            {"id": 4, "start": None, "end": None, "a": 10.0, "e": 2100000.0},
            {"id": 5, "start": None, "end": None, "a": 10.0, "e": 2100000.0},
            {"id": 6, "start": None, "end": None, "a": 10.0, "e": 2100000.0},
            {"id": 7, "start": None, "end": None, "a": 10.0, "e": 2100000.0},
            {"id": 8, "start": None, "end": None, "a": 10.0, "e": 2100000.0},
        ]
        n_free = 9  # Up to 9 free DOFs for 6 nodes isostatic (12 DOFs - 3 reactions)
        n_rest = 3  # Exactly 3 reactions for isostatic
        sheet_hero_title = "PLANTILLA MANUAL - CASO ISOSTÁTICO (HASTA 6 NODOS Y 8 BARRAS - CÁTEDRA ING. ULIANOV)"
    else:
        # Default: 4 nodes active (Cruz de San Andrés), nodes 5-6 ready for input
        nodes_data = [
            {"id": 1, "x": 0.0,   "y": 0.0,   "support": "Fijo",  "px": 0.0,    "py": 0.0},
            {"id": 2, "x": 400.0, "y": 0.0,   "support": "Fijo",  "px": 0.0,    "py": 0.0},
            {"id": 3, "x": 400.0, "y": 300.0, "support": "Libre", "px": 3500.0, "py": -5000.0},
            {"id": 4, "x": 0.0,   "y": 300.0, "support": "Libre", "px": 3500.0, "py": 0.0},
            {"id": 5, "x": None,  "y": None,  "support": "Libre", "px": None,   "py": None},
            {"id": 6, "x": None,  "y": None,  "support": "Libre", "px": None,   "py": None},
        ]
        bars_data = [
            {"id": 1, "start": 1, "end": 2, "a": 12.0, "e": 2100000.0},
            {"id": 2, "start": 2, "end": 3, "a": 12.0, "e": 2100000.0},
            {"id": 3, "start": 3, "end": 4, "a": 12.0, "e": 2100000.0},
            {"id": 4, "start": 4, "end": 1, "a": 12.0, "e": 2100000.0},
            {"id": 5, "start": 1, "end": 3, "a": 12.0, "e": 2100000.0},
            {"id": 6, "start": 2, "end": 4, "a": 12.0, "e": 2100000.0},
            {"id": 7, "start": None, "end": None, "a": 12.0, "e": 2100000.0},
            {"id": 8, "start": None, "end": None, "a": 12.0, "e": 2100000.0},
        ]
        n_free = 8  # Up to 8 free DOFs for 4-6 nodes hyperstatic (12 DOFs - 4 reactions)
        n_rest = 4  # 4 reactions (GH=1..2)
        sheet_hero_title = "PLANTILLA MANUAL - CASO HIPERESTÁTICO (HASTA 6 NODOS Y 8 BARRAS - CÁTEDRA ING. ULIANOV)"

    dofs_list = [f"U{(i//2)+1}{'X' if i%2==0 else 'Y'}" for i in range(dofs_total)]

    # Layout column widths
    pedag_c_start = 26 # P to W used for Nodal Table; X, Y spacers; Z to AE for pedagogical cards
    pedag_c_end = pedag_c_start + 5

    for col_idx in range(1, 15):
        col_let = get_column_letter(col_idx)
        ws.column_dimensions[col_let].width = 13
    ws.column_dimensions['A'].width = 11
    ws.column_dimensions['O'].width = 3
    ws.column_dimensions['P'].width = 7  # Nudo ID
    ws.column_dimensions['Q'].width = 11 # X
    ws.column_dimensions['R'].width = 11 # Y
    ws.column_dimensions['S'].width = 14 # Tipo Apoyo
    ws.column_dimensions['T'].width = 6  # Rx
    ws.column_dimensions['U'].width = 6  # Ry
    ws.column_dimensions['V'].width = 12 # Px
    ws.column_dimensions['W'].width = 12 # Py
    ws.column_dimensions['X'].width = 3  # spacer
    ws.column_dimensions['Y'].width = 3  # spacer
    for c_p in range(pedag_c_start, pedag_c_end + 1):
        ws.column_dimensions[get_column_letter(c_p)].width = 18

    # Row 1-2: Hero Header
    ws.merge_cells("A1:N1")
    ws["A1"] = sheet_hero_title
    ws["A1"].font = f_hero
    ws["A1"].fill = fill_hero
    ws["A1"].alignment = al_center
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:N2")
    ws["A2"] = "Cátedra: Análisis Estructural II | Ing. Ulianov Cuba Valencia | Plantilla 100% Automatizada con Capacidad hasta 6 Nodos y 8 Barras"
    ws["A2"].font = f_hero_sub
    ws["A2"].fill = fill_hero
    ws["A2"].alignment = al_center
    ws.row_dimensions[2].height = 18

    # Table 1: Headers (Rows 4-5)
    t1_headers = [
        (4, 5, 2, 2, "BARRA"),
        (4, 4, 3, 4, "INICIO (cm)"),
        (5, 5, 3, 3, "X"),
        (5, 5, 4, 4, "Y"),
        (4, 4, 5, 6, "FIN (cm)"),
        (5, 5, 5, 5, "X"),
        (5, 5, 6, 6, "Y"),
        (4, 5, 7, 7, "DX (cm)"),
        (4, 5, 8, 8, "DY (cm)"),
        (4, 5, 9, 9, "L (cm)"),
        (4, 5, 10, 10, "Cx"),
        (4, 5, 11, 11, "Cy"),
        (4, 5, 12, 12, "A (cm²)"),
        (4, 5, 13, 13, "E (Kgf/cm²)"),
        (4, 5, 14, 14, "AE/L (Kgf/cm)"),
    ]
    for r1, r2, c1, c2, txt in t1_headers:
        if r1 != r2 or c1 != c2:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(r1, c1)
        cell.value = txt
        cell.alignment = al_center

    apply_styling(ws, 4, 4, 2, 14, fill=fill_sec, font=f_tbl_hdr, align=al_center)
    apply_styling(ws, 5, 5, 2, 14, fill=fill_tbl_sub, font=f_tbl_sub, align=al_center)
    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 18

    # Table of Nodal Coordinates & Support Conditions in Cols P to W (Rows 4-11 for 6 nodes)
    ws.merge_cells("P4:W4")
    ws["P4"] = "COORDENADAS, CONDICIONES DE APOYO Y CARGAS NODALES (HASTA 6 NODOS)"
    ws["P4"].font = f_sec_hdr
    ws["P4"].fill = fill_sec_teal
    ws["P4"].alignment = al_center

    node_hdrs = ["NUDO", "X (cm)", "Y (cm)", "TIPO APOYO", "Rx", "Ry", "Px (Kgf)", "Py (Kgf)"]
    for idx_nh, nh in enumerate(node_hdrs):
        c = 16 + idx_nh
        ws.cell(5, c, nh)
    apply_styling(ws, 5, 5, 16, 23, fill=fill_tbl_sub, font=f_tbl_sub, align=al_center)

    nodal_table_range = f"$P$6:$W${5 + MAX_NODES}"

    for idx_n, nd in enumerate(nodes_data, start=1):
        r_n = 5 + idx_n
        ws.cell(r_n, 16, nd["id"])
        ws.cell(r_n, 17, nd["x"])
        ws.cell(r_n, 18, nd["y"])
        ws.cell(r_n, 19, nd["support"])
        ws.cell(r_n, 20, f'=IF(Q{r_n}="","", IF(LEFT(UPPER(TRIM(S{r_n})),4)="FIJO", 1, IF(OR(UPPER(TRIM(S{r_n}))="MOVIL X", UPPER(TRIM(S{r_n}))="MÓVIL X", UPPER(TRIM(S{r_n}))="RODILLO X"), 1, 0)))')
        ws.cell(r_n, 21, f'=IF(Q{r_n}="","", IF(LEFT(UPPER(TRIM(S{r_n})),4)="FIJO", 1, IF(OR(UPPER(TRIM(S{r_n}))="MOVIL Y", UPPER(TRIM(S{r_n}))="MÓVIL Y", UPPER(TRIM(S{r_n}))="RODILLO Y", UPPER(TRIM(S{r_n}))="MOVIL", UPPER(TRIM(S{r_n}))="MÓVIL"), 1, 0)))')
        ws.cell(r_n, 22, nd["px"])
        ws.cell(r_n, 23, nd["py"])

        apply_styling(ws, r_n, r_n, 16, 16, fill=fill_card, font=f_bold, align=al_center)
        apply_styling(ws, r_n, r_n, 17, 18, fill=fill_input, font=f_input, align=al_right, num_fmt="#,##0.00")
        apply_styling(ws, r_n, r_n, 19, 19, fill=fill_input, font=f_input, align=al_center)
        apply_styling(ws, r_n, r_n, 20, 21, fill=fill_card, font=f_bold, align=al_center)
        apply_styling(ws, r_n, r_n, 22, 23, fill=fill_input, font=f_input, align=al_right, num_fmt="#,##0.00")
        ws.row_dimensions[r_n].height = 19

    # Table 2: Conectividad y Grados de Libertad (Rows 15 to 24)
    t2_title_row = 15
    ws.merge_cells(f"B{t2_title_row}:H{t2_title_row}")
    ws.cell(t2_title_row, 2, "CONECTIVIDAD DE BARRAS Y GRADOS DE LIBERTAD ASOCIADOS:")
    ws.cell(t2_title_row, 2).font = f_sec_hdr
    ws.cell(t2_title_row, 2).fill = fill_sec
    ws.row_dimensions[t2_title_row].height = 20

    t2_hdr_row = t2_title_row + 1
    t2_headers = [
        (t2_hdr_row, t2_hdr_row, 2, 2, "BARRA"),
        (t2_hdr_row, t2_hdr_row, 3, 3, "NUDO INICIO"),
        (t2_hdr_row, t2_hdr_row, 4, 4, "u_ini_x"),
        (t2_hdr_row, t2_hdr_row, 5, 5, "u_ini_y"),
        (t2_hdr_row, t2_hdr_row, 6, 6, "NUDO FIN"),
        (t2_hdr_row, t2_hdr_row, 7, 7, "u_fin_x"),
        (t2_hdr_row, t2_hdr_row, 8, 8, "u_fin_y"),
    ]
    for r1, r2, c1, c2, txt in t2_headers:
        ws.cell(r1, c1, txt)
        ws.cell(r1, c1).alignment = al_center
    apply_styling(ws, t2_hdr_row, t2_hdr_row, 2, 8, fill=fill_sec, font=f_tbl_hdr, align=al_center)
    ws.row_dimensions[t2_hdr_row].height = 18

    t2_data_start = t2_hdr_row + 1

    # Populate Table 2 (Rows 17 to 24 for 8 bars)
    for m in range(1, MAX_BARS + 1):
        r2 = t2_data_start + m - 1
        ws.row_dimensions[r2].height = 19
        br = bars_data[m - 1]
        ni = br["start"]
        nf = br["end"]

        bar_col_info = BAR_PASTEL_PALETTE[(m - 1) % len(BAR_PASTEL_PALETTE)]
        fill_bar_tracer = PatternFill("solid", fgColor=bar_col_info["bg"])

        ws.cell(r2, 2, m)
        if ni is not None:
            ws.cell(r2, 3, ni)
            ws.cell(r2, 6, nf)
        else:
            ws.cell(r2, 3, None)
            ws.cell(r2, 6, None)

        ws.cell(r2, 4, f'=IF(C{r2}="","", "U" & C{r2} & "X")')
        ws.cell(r2, 5, f'=IF(C{r2}="","", "U" & C{r2} & "Y")')
        ws.cell(r2, 7, f'=IF(F{r2}="","", "U" & F{r2} & "X")')
        ws.cell(r2, 8, f'=IF(F{r2}="","", "U" & F{r2} & "Y")')

        apply_styling(ws, r2, r2, 2, 2, fill=fill_bar_tracer, font=f_bold, align=al_center)
        apply_styling(ws, r2, r2, 3, 3, fill=fill_input, font=f_input, align=al_center, num_fmt="0")
        apply_styling(ws, r2, r2, 4, 5, fill=fill_card, font=f_bold, align=al_center)
        apply_styling(ws, r2, r2, 6, 6, fill=fill_input, font=f_input, align=al_center, num_fmt="0")
        apply_styling(ws, r2, r2, 7, 8, fill=fill_card, font=f_bold, align=al_center)

    # Populate Table 1 Data Rows (Rows 6 to 13 for 8 bars)
    for m in range(1, MAX_BARS + 1):
        r1 = 5 + m
        r2 = t2_data_start + m - 1
        ws.row_dimensions[r1].height = 19
        br = bars_data[m - 1]

        bar_col_info = BAR_PASTEL_PALETTE[(m - 1) % len(BAR_PASTEL_PALETTE)]
        fill_bar_tracer = PatternFill("solid", fgColor=bar_col_info["bg"])

        ws.cell(r1, 2, m)
        apply_styling(ws, r1, r1, 2, 2, fill=fill_bar_tracer, font=f_bold, align=al_center)

        ws.cell(r1, 3, f'=IF(C{r2}="","", IFERROR(VLOOKUP(C{r2}, {nodal_table_range}, 2, FALSE), ""))')
        ws.cell(r1, 4, f'=IF(C{r2}="","", IFERROR(VLOOKUP(C{r2}, {nodal_table_range}, 3, FALSE), ""))')
        ws.cell(r1, 5, f'=IF(F{r2}="","", IFERROR(VLOOKUP(F{r2}, {nodal_table_range}, 2, FALSE), ""))')
        ws.cell(r1, 6, f'=IF(F{r2}="","", IFERROR(VLOOKUP(F{r2}, {nodal_table_range}, 3, FALSE), ""))')
        apply_styling(ws, r1, r1, 3, 6, fill=fill_input, font=f_input, align=al_right, num_fmt="#,##0.00")

        ws.cell(r1, 7, f'=IF(OR(C{r1}="",E{r1}=""), "", E{r1}-C{r1})')
        ws.cell(r1, 8, f'=IF(OR(D{r1}="",F{r1}=""), "", F{r1}-D{r1})')
        ws.cell(r1, 9, f'=IF(OR(G{r1}="",H{r1}=""), "", SQRT(G{r1}^2+H{r1}^2))')
        ws.cell(r1, 10, f'=IF(OR(I{r1}="",I{r1}=0), "", G{r1}/I{r1})')
        ws.cell(r1, 11, f'=IF(OR(I{r1}="",I{r1}=0), "", H{r1}/I{r1})')
        apply_styling(ws, r1, r1, 7, 8, fill=fill_card, font=f_regular, align=al_right, num_fmt="#,##0.00")
        apply_styling(ws, r1, r1, 9, 9, fill=fill_card, font=f_bold, align=al_right, num_fmt="#,##0.00")
        apply_styling(ws, r1, r1, 10, 11, fill=fill_card, font=f_regular, align=al_right, num_fmt="0.0000")

        ws.cell(r1, 12, f'=IF(C{r2}="","", {br["a"]})')
        ws.cell(r1, 13, f'=IF(C{r2}="","", {br["e"]})')
        apply_styling(ws, r1, r1, 12, 12, fill=fill_input, font=f_input, align=al_right, num_fmt="#,##0.00")
        apply_styling(ws, r1, r1, 13, 13, fill=fill_input, font=f_input, align=al_right, num_fmt="#,##0")

        ws.cell(r1, 14, f'=IF(OR(I{r1}="",I{r1}=0,L{r1}="",M{r1}=""), 0, (L{r1}*M{r1})/I{r1})')
        apply_styling(ws, r1, r1, 14, 14, fill=fill_card, font=f_bold, align=al_right, num_fmt="#,##0.00")

    # Section 3: Matrices Locales (4x4) y Expandidas (12x12) para las 8 BARRAS
    curr_r = t2_data_start + MAX_BARS + 2
    factors = [
        ("J", "J", 1), ("J", "K", 1), ("J", "J", -1), ("J", "K", -1),
        ("J", "K", 1), ("K", "K", 1), ("J", "K", -1), ("K", "K", -1),
        ("J", "J", -1), ("J", "K", -1), ("J", "J", 1), ("J", "K", 1),
        ("J", "K", -1), ("K", "K", -1), ("J", "K", 1), ("K", "K", 1),
    ]
    exp_matrix_start_rows = []

    for m in range(1, MAX_BARS + 1):
        t1_row = 5 + m
        t2_row = t2_data_start + (m - 1)
        bar_col_info = BAR_PASTEL_PALETTE[(m - 1) % len(BAR_PASTEL_PALETTE)]
        fill_local = PatternFill("solid", fgColor=bar_col_info["bg"])
        fill_hdr_bar = PatternFill("solid", fgColor=bar_col_info["bg"])

        # 4x4 Headers
        hdr_4x4_row = curr_r
        ws.row_dimensions[hdr_4x4_row].height = 19
        ws.cell(hdr_4x4_row, 1, f"k{m}=")
        ws.cell(hdr_4x4_row, 1).font = f_matrix_sym
        ws.cell(hdr_4x4_row, 1).alignment = al_center

        ws.cell(hdr_4x4_row, 2, f'=IF(D{t2_row}="","-",D{t2_row})')
        ws.cell(hdr_4x4_row, 3, f'=IF(E{t2_row}="","-",E{t2_row})')
        ws.cell(hdr_4x4_row, 4, f'=IF(G{t2_row}="","-",G{t2_row})')
        ws.cell(hdr_4x4_row, 5, f'=IF(H{t2_row}="","-",H{t2_row})')
        apply_styling(ws, hdr_4x4_row, hdr_4x4_row, 2, 5, fill=fill_hdr_bar, font=f_bold, align=al_center)

        # 4x4 Content
        mat_4x4_row = hdr_4x4_row + 1
        dof_cols_t2 = ["D", "E", "G", "H"]
        for i_row in range(4):
            r = mat_4x4_row + i_row
            ws.row_dimensions[r].height = 18
            ws.cell(r, 6, f'=IF({dof_cols_t2[i_row]}{t2_row}="","-",{dof_cols_t2[i_row]}{t2_row})')
            apply_styling(ws, r, r, 6, 6, fill=fill_hdr_bar, font=f_bold, align=al_center)

            for j_col in range(4):
                f_idx = i_row * 4 + j_col
                c1, c2, sign = factors[f_idx]
                sign_str = "-" if sign == -1 else ""
                ws.cell(r, 2 + j_col, f'=IF(N{t1_row}=0, 0, {sign_str}N{t1_row}*{c1}{t1_row}*{c2}{t1_row})')
                apply_styling(ws, r, r, 2 + j_col, 2 + j_col, fill=fill_local, font=f_regular, align=al_right, num_fmt="#,##0.00")

        # Expanded Matrix (K_m: 12x12)
        hdr_exp_row = mat_4x4_row + 4 + 1
        ws.row_dimensions[hdr_exp_row].height = 19
        for c_idx, dof in enumerate(dofs_list, start=2):
            ws.cell(hdr_exp_row, c_idx, dof)
        apply_styling(ws, hdr_exp_row, hdr_exp_row, 2, 1 + dofs_total, fill=fill_sec, font=f_tbl_hdr, align=al_center)

        mat_exp_row = hdr_exp_row + 1
        exp_matrix_start_rows.append(mat_exp_row)

        ws.cell(mat_exp_row + (dofs_total // 2) - 1, 1, f"K{m}=")
        ws.cell(mat_exp_row + (dofs_total // 2) - 1, 1).font = f_matrix_sym
        ws.cell(mat_exp_row + (dofs_total // 2) - 1, 1).alignment = al_center

        lbl_col_idx = 2 + dofs_total # Col N (14)
        lbl_col_let = get_column_letter(lbl_col_idx)

        for i_row, dof in enumerate(dofs_list):
            curr_exp_r = mat_exp_row + i_row
            ws.row_dimensions[curr_exp_r].height = 19
            ws.cell(curr_exp_r, lbl_col_idx, dof)
            for i_col in range(2, 2 + dofs_total):
                col_let = get_column_letter(i_col)
                formula = f'=IFERROR(INDEX($B${mat_4x4_row}:$E${mat_4x4_row+3}, MATCH(${lbl_col_let}{curr_exp_r}, $B${hdr_4x4_row}:$E${hdr_4x4_row}, 0), MATCH({col_let}${hdr_exp_row}, $B${hdr_4x4_row}:$E${hdr_4x4_row}, 0)), 0)'
                ws.cell(curr_exp_r, i_col, formula)
                apply_styling(ws, curr_exp_r, curr_exp_r, i_col, i_col, fill=fill_white, font=f_regular, align=al_right, num_fmt="#,##0.00")

            apply_styling(ws, curr_exp_r, curr_exp_r, lbl_col_idx, lbl_col_idx, fill=fill_hdr_bar, font=f_bold, align=al_center)

        # COLOR TRACKING: Conditional formatting rule so non-zero cells in K_m take the bar's pastel color!
        rule_bar = CellIsRule(operator='notEqual', formula=['0'], fill=fill_local, font=f_bold)
        k_range_str = f"B{mat_exp_row}:{get_column_letter(1 + dofs_total)}{mat_exp_row + dofs_total - 1}"
        ws.conditional_formatting.add(k_range_str, rule_bar)

        curr_r = mat_exp_row + dofs_total + 2

    # Section 4: Karmadura (12x12) Sum of all 8 bars with uniform pastel green diagonal
    hdr_karm_row = curr_r
    karm_end_col_let = get_column_letter(1 + dofs_total)
    lbl_col_idx = 2 + dofs_total
    lbl_col_let = get_column_letter(lbl_col_idx)
    ws.row_dimensions[hdr_karm_row].height = 20
    for c_idx, dof in enumerate(dofs_list, start=2):
        ws.cell(hdr_karm_row, c_idx, dof)
    apply_styling(ws, hdr_karm_row, hdr_karm_row, 2, 1 + dofs_total, fill=fill_sec, font=f_tbl_hdr, align=al_center)

    mat_karm_row = hdr_karm_row + 1
    mat_karm_end = mat_karm_row + dofs_total - 1
    ws.cell(mat_karm_row + (dofs_total // 2) - 1, 1, "Karmadura=")
    ws.cell(mat_karm_row + (dofs_total // 2) - 1, 1).font = f_matrix_sym
    ws.cell(mat_karm_row + (dofs_total // 2) - 1, 1).alignment = al_center

    for i_r, dof in enumerate(dofs_list):
        curr_k_r = mat_karm_row + i_r
        ws.row_dimensions[curr_k_r].height = 20
        ws.cell(curr_k_r, lbl_col_idx, dof)
        apply_styling(ws, curr_k_r, curr_k_r, lbl_col_idx, lbl_col_idx, fill=fill_tbl_sub, font=f_bold, align=al_center)

        for i_c in range(2, 2 + dofs_total):
            c_let = get_column_letter(i_c)
            formula = "=" + "+".join(f"{c_let}{exp_r + i_r}" for exp_r in exp_matrix_start_rows)
            ws.cell(curr_k_r, i_c, formula)

            node_r = i_r // 2
            node_c = (i_c - 2) // 2
            if node_r == node_c:
                cell_fill = fill_karm_diag
                cell_font = f_bold
            else:
                cell_fill = fill_white
                cell_font = f_regular

            cell_border = border_diag_med if (i_r == (i_c - 2)) else border_thin
            apply_styling(ws, curr_k_r, curr_k_r, i_c, i_c, fill=cell_fill, font=cell_font, align=al_right, border=cell_border, num_fmt="#,##0.00")

    # 1. MATRIZ DE CARGAS {C} (12 DOFs)
    cargas_hdr_row = mat_karm_row + dofs_total + 2
    ws.cell(cargas_hdr_row, 1, "Matriz de cargas:")
    ws.cell(cargas_hdr_row, 1).font = f_prof_title_no_ul
    ws.row_dimensions[cargas_hdr_row].height = 20

    cargas_start_row = cargas_hdr_row + 1
    cargas_end_row = cargas_start_row + dofs_total - 1
    mid_carga_row = cargas_start_row + (dofs_total // 2) - 1

    ws.cell(mid_carga_row, 1, "C=")
    apply_styling(ws, mid_carga_row, mid_carga_row, 1, 1, fill=fill_prof_green, font=f_bold, align=al_center)

    for idx_d in range(dofs_total):
        r = cargas_start_row + idx_d
        ws.row_dimensions[r].height = 19
        node_num = (idx_d // 2) + 1
        is_x = (idx_d % 2 == 0)
        dof_var_c = f"C{node_num}{'X' if is_x else 'Y'}"
        dof_lbl_u = f"U{node_num}{'X' if is_x else 'Y'}"

        ws.cell(r, 2, dof_var_c)
        ws.cell(r, 3, dof_lbl_u)
        apply_styling(ws, r, r, 2, 3, fill=fill_white, font=f_bold, align=al_center)

        if r == mid_carga_row:
            ws.cell(r, 4, "=")
            ws.cell(r, 4).alignment = al_center

        rest_col_idx = 5 if is_x else 6
        load_col_idx = 7 if is_x else 8
        c_formula = (
            f'=IF(IFERROR(VLOOKUP({node_num}, {nodal_table_range}, 2, FALSE), "")="", 0, '
            f'IF(VLOOKUP({node_num}, {nodal_table_range}, {rest_col_idx}, FALSE)=1, B{r}, '
            f'VLOOKUP({node_num}, {nodal_table_range}, {load_col_idx}, FALSE)))'
        )
        ws.cell(r, 5, c_formula)
        apply_styling(ws, r, r, 5, 5, fill=fill_peach, font=f_bold, align=al_right, num_fmt="#,##0.00")

    # 2. MATRIZ DE DESPLAZAMIENTOS {D} (12 DOFs)
    disp_hdr_row = cargas_end_row + 2
    ws.cell(disp_hdr_row, 1, "Matriz de desplazamientos:")
    ws.cell(disp_hdr_row, 1).font = f_prof_title_no_ul
    ws.row_dimensions[disp_hdr_row].height = 20

    disp_start_row = disp_hdr_row + 1
    disp_end_row = disp_start_row + dofs_total - 1
    mid_disp_row = disp_start_row + (dofs_total // 2) - 1

    ws.cell(mid_disp_row, 1, "D=")
    apply_styling(ws, mid_disp_row, mid_disp_row, 1, 1, fill=fill_prof_green, font=f_bold, align=al_center)

    for idx_d in range(dofs_total):
        r = disp_start_row + idx_d
        src_c_r = cargas_start_row + idx_d
        ws.row_dimensions[r].height = 19
        node_num = (idx_d // 2) + 1
        is_x = (idx_d % 2 == 0)
        dof_var_d = f"D{node_num}{'X' if is_x else 'Y'}"
        dof_lbl_u = f"U{node_num}{'X' if is_x else 'Y'}"

        ws.cell(r, 2, dof_var_d)
        ws.cell(r, 3, dof_lbl_u)
        apply_styling(ws, r, r, 2, 3, fill=fill_white, font=f_bold, align=al_center)

        if r == mid_disp_row:
            ws.cell(r, 4, "=")
            ws.cell(r, 4).alignment = al_center

        rest_col_idx = 5 if is_x else 6
        d_formula = (
            f'=IF(IFERROR(VLOOKUP({node_num}, {nodal_table_range}, 2, FALSE), "")="", 0, '
            f'IF(VLOOKUP({node_num}, {nodal_table_range}, {rest_col_idx}, FALSE)=1, 0, B{r}))'
        )
        ws.cell(r, 5, d_formula)
        apply_styling(ws, r, r, 5, 5, fill=fill_peach, font=f_bold, align=al_center)

        # Helper Col F: Running count of FREE DOFs
        if idx_d == 0:
            ws.cell(r, 6, f'=IF(ISTEXT(E{r}), 1, "")')
        else:
            ws.cell(r, 6, f'=IF(ISTEXT(E{r}), MAX(F${disp_start_row}:F{r-1}) + 1, "")')
        apply_styling(ws, r, r, 6, 6, fill=fill_card, font=f_note, align=al_center)

        # Helper Col G: Running count of RESTRAINED DOFs
        if idx_d == 0:
            ws.cell(r, 7, f'=IF(ISTEXT(E{src_c_r}), 1, "")')
        else:
            ws.cell(r, 7, f'=IF(ISTEXT(E{src_c_r}), MAX(G${disp_start_row}:G{r-1}) + 1, "")')
        apply_styling(ws, r, r, 7, 7, fill=fill_card, font=f_note, align=al_center)

    # Pedagogical Card 1
    card1_lines = [
        "• Columna B: Variable del vector (C1X, C1Y / D1X, D1Y) hasta 6 nodos (12 GDL).",
        "• Columna C: Grado de libertad global asociado (U1X, U1Y...).",
        "• Columna E (Cargas): Si hay apoyo, C=Cnx (reacción incógnita). Si está libre, C=Px (carga externa).",
        "• Columna E (Desplazamientos): Si hay apoyo, D=0 (restringido). Si está libre, D=Dnx (desplazamiento incógnita).",
        "• Columnas F y G (Auxiliares): Conteo dinámico de grados libres (D≠0) y apoyos (D=0) para extraer K11 y K21."
    ]
    write_pedag_card(ws, cargas_hdr_row, disp_end_row, pedag_c_start, pedag_c_end,
                      "1. MATRIZ DE CARGAS Y DESPLAZAMIENTOS ({C} Y {D})", card1_lines)

    # 3. ECUACION MATRICIAL: C = Karmadura*D (12x12)
    sys_hdr_row = disp_end_row + 2
    ws.cell(sys_hdr_row, 1, "Ecuacion matricial:")
    ws.cell(sys_hdr_row, 1).font = f_prof_title
    ws.row_dimensions[sys_hdr_row].height = 22

    ws.merge_cells(start_row=sys_hdr_row, start_column=3, end_row=sys_hdr_row, end_column=4)
    ws.cell(sys_hdr_row, 3, "C = Karmadura*D")
    apply_styling(ws, sys_hdr_row, sys_hdr_row, 3, 4, fill=fill_yellow_banner, font=f_yellow_banner, align=al_center)

    ws.cell(sys_hdr_row, 5, "ELIMINO FILA Y COLUMNA DE LOS QUE NO CORRESPONDE PARA HALLAR LOS DESPLAZAMIENTOS")
    ws.cell(sys_hdr_row, 5).font = f_prof_note
    ws.cell(sys_hdr_row + 1, 5, "ELIMINO COLUMNA SEGÚN LAS REACCIONES EN ORDEN SIN CONSIDERAR LAS FILAS ANULADAS")
    ws.cell(sys_hdr_row + 1, 5).font = f_prof_note
    ws.row_dimensions[sys_hdr_row + 1].height = 18

    sys_data_start = sys_hdr_row + 3
    ws.cell(sys_data_start - 1, 1, "C")
    ws.cell(sys_data_start - 1, 1).font = f_bold
    ws.cell(sys_data_start - 1, 1).alignment = al_center

    ws.cell(sys_data_start - 1, 3, "K")
    ws.cell(sys_data_start - 1, 3).font = f_bold
    ws.cell(sys_data_start - 1, 3).alignment = al_center

    disp_col_idx = 3 + dofs_total + 1
    ws.cell(sys_data_start - 1, disp_col_idx, "D")
    ws.cell(sys_data_start - 1, disp_col_idx).font = f_bold
    ws.cell(sys_data_start - 1, disp_col_idx).alignment = al_center

    for idx_d in range(dofs_total):
        r = sys_data_start + idx_d
        ws.row_dimensions[r].height = 20
        src_cargas_r = cargas_start_row + idx_d
        src_disp_r = disp_start_row + idx_d

        ws.cell(r, 1, f"=E{src_cargas_r}")
        apply_styling(ws, r, r, 1, 1, fill=fill_white, font=f_bold, align=al_right, num_fmt="#,##0.00")

        if idx_d == (dofs_total // 2):
            ws.cell(r, 2, "=")
            ws.cell(r, 2).alignment = al_center

        for col_k in range(dofs_total):
            c_dest = 3 + col_k
            c_src = 2 + col_k
            src_let = get_column_letter(c_src)
            src_row = mat_karm_row + idx_d
            ws.cell(r, c_dest, f"={src_let}{src_row}")
            apply_styling(ws, r, r, c_dest, c_dest, fill=fill_white, font=f_regular, align=al_right, num_fmt="#,##0.00")

        ws.cell(r, 3 + dofs_total, "·")
        ws.cell(r, 3 + dofs_total).alignment = al_center

        ws.cell(r, disp_col_idx, f"=E{src_disp_r}")
        apply_styling(ws, r, r, disp_col_idx, disp_col_idx, fill=fill_white, font=f_bold, align=al_center)

    # Pedagogical Card 2
    card2_lines = [
        "• Fundamento: Ley de Hooke generalizada para armaduras planas {C} = [K_global] · {D}.",
        "• Apunte del Docente: 'ELIMINO FILA Y COLUMNA DE LOS QUE NO CORRESPONDE PARA HALLAR LOS DESPLAZAMIENTOS'.",
        "• Partición de Rigidez:",
        "    [ C_libres ] = [ K11   K12 ] · [ D_libres ]",
        "    [ R_apoyos ]   [ K21   K22 ]   [    0     ]",
        "• K11 (verde pastel): Rigideces directas para {D_libres} = [K11]⁻¹ · {C_libres}.",
        "• K21 (lavanda pastel): Rigideces de acoplamiento para reacciones {R} = [K21] · {D_libres}."
    ]
    write_pedag_card(ws, sys_hdr_row, sys_data_start + dofs_total - 1, pedag_c_start, pedag_c_end,
                      "2. ECUACIÓN MATRICIAL GLOBAL: {C} = [K] · {D}", card2_lines)

    # 4. ECUACION MATRICIAL REDUCIDA PARA EL CALCULO DE DESPLAZAMIENTOS (K11: n_free x n_free)
    k11_hdr_row = sys_data_start + dofs_total + 2
    ws.cell(k11_hdr_row, 1, "Ecuacion matricial reducida para el calculo de desplazamientos:")
    ws.cell(k11_hdr_row, 1).font = f_prof_title
    ws.row_dimensions[k11_hdr_row].height = 22

    k11_col_dofs_row = k11_hdr_row + 1
    ws.row_dimensions[k11_col_dofs_row].height = 18

    for j_f in range(n_free):
        c_dest = 3 + j_f
        k = j_f + 1
        ws.cell(k11_col_dofs_row, c_dest, f'=IFERROR(INDEX(C${disp_start_row}:C${disp_end_row}, MATCH({k}, F${disp_start_row}:F${disp_end_row}, 0)), "-")')
        apply_styling(ws, k11_col_dofs_row, k11_col_dofs_row, c_dest, c_dest, fill=fill_prof_green, font=f_bold, align=al_center)

    k11_data_start = k11_col_dofs_row + 1
    k11_end_row = k11_data_start + n_free - 1
    mid_k11_row = k11_data_start + (n_free // 2)

    k11_helper_col = 3 + n_free + 2
    k11_helper_col_let = get_column_letter(k11_helper_col)
    d_lbl_col = 3 + n_free + 1

    for i_f in range(n_free):
        r = k11_data_start + i_f
        ws.row_dimensions[r].height = 20
        k = i_f + 1

        ws.cell(r, k11_helper_col, f'=IFERROR(INDEX(C${disp_start_row}:C${disp_end_row}, MATCH({k}, F${disp_start_row}:F${disp_end_row}, 0)), "-")')
        apply_styling(ws, r, r, k11_helper_col, k11_helper_col, fill=fill_card, font=f_note, align=al_center)

        ws.cell(r, d_lbl_col, f'=IFERROR(INDEX(B${disp_start_row}:B${disp_end_row}, MATCH({k}, F${disp_start_row}:F${disp_end_row}, 0)), "-")')
        apply_styling(ws, r, r, d_lbl_col, d_lbl_col, fill=fill_prof_green, font=f_bold, align=al_center)

        ws.cell(r, 1, f'=IF(${k11_helper_col_let}{r}="-", 0, IFERROR(INDEX(E${cargas_start_row}:E${cargas_end_row}, MATCH(${k11_helper_col_let}{r}, C${cargas_start_row}:C${cargas_end_row}, 0)), 0))')
        apply_styling(ws, r, r, 1, 1, fill=fill_prof_green, font=f_bold, align=al_right, num_fmt="#,##0.00")

        if r == mid_k11_row:
            ws.cell(r, 2, "=")
            ws.cell(r, 2).alignment = al_center

        for j_f in range(n_free):
            c_dest = 3 + j_f
            col_dest_let = get_column_letter(c_dest)
            inact_val = "1" if i_f == j_f else "0"
            formula = (
                f'=IF(OR(${k11_helper_col_let}{r}="-", {col_dest_let}${k11_col_dofs_row}="-"), {inact_val}, '
                f'IF(OR('
                f'IFERROR(INDEX($B${mat_karm_row}:${karm_end_col_let}${mat_karm_end}, MATCH(${k11_helper_col_let}{r}, ${lbl_col_let}${mat_karm_row}:${lbl_col_let}${mat_karm_end}, 0), MATCH(${k11_helper_col_let}{r}, $B${hdr_karm_row}:${karm_end_col_let}${hdr_karm_row}, 0)), 0)=0, '
                f'IFERROR(INDEX($B${mat_karm_row}:${karm_end_col_let}${mat_karm_end}, MATCH({col_dest_let}${k11_col_dofs_row}, ${lbl_col_let}${mat_karm_row}:${lbl_col_let}${mat_karm_end}, 0), MATCH({col_dest_let}${k11_col_dofs_row}, $B${hdr_karm_row}:${karm_end_col_let}${hdr_karm_row}, 0)), 0)=0), '
                f'{inact_val}, '
                f'IFERROR(INDEX($B${mat_karm_row}:${karm_end_col_let}${mat_karm_end}, MATCH(${k11_helper_col_let}{r}, ${lbl_col_let}${mat_karm_row}:${lbl_col_let}${mat_karm_end}, 0), MATCH({col_dest_let}${k11_col_dofs_row}, $B${hdr_karm_row}:${karm_end_col_let}${hdr_karm_row}, 0)), 0)))'
            )
            ws.cell(r, c_dest, formula)
            apply_styling(ws, r, r, c_dest, c_dest, fill=fill_prof_green, font=f_bold, align=al_right, num_fmt="#,##0.00")

        ws.cell(r, 3 + n_free, "·")
        ws.cell(r, 3 + n_free).alignment = al_center

    # 5. D=MINVERSA(K)*C
    inv_banner_row = k11_end_row + 2
    banner_end_col = max(3 + n_free - 1, 4)
    ws.merge_cells(start_row=inv_banner_row, start_column=3, end_row=inv_banner_row, end_column=banner_end_col)
    ws.cell(inv_banner_row, 3, "D=MINVERSA(K)*C")
    apply_styling(ws, inv_banner_row, inv_banner_row, 3, banner_end_col, fill=fill_yellow_banner, font=f_yellow_banner, align=al_center)
    ws.row_dimensions[inv_banner_row].height = 20

    inv_hdr_row = inv_banner_row + 1
    ws.cell(inv_hdr_row, 1, "D")
    ws.cell(inv_hdr_row, 1).font = f_bold
    ws.cell(inv_hdr_row, 1).alignment = al_center

    if n_free > 1:
        ws.merge_cells(start_row=inv_hdr_row, start_column=3, end_row=inv_hdr_row, end_column=3 + n_free - 1)
    ws.cell(inv_hdr_row, 3, "MIN")
    ws.cell(inv_hdr_row, 3).font = f_bold
    ws.cell(inv_hdr_row, 3).alignment = al_center

    c_vec_col_idx = 3 + n_free + 1
    c_vec_col_let = get_column_letter(c_vec_col_idx)
    ws.cell(inv_hdr_row, c_vec_col_idx, "C")
    ws.cell(inv_hdr_row, c_vec_col_idx).font = f_bold
    ws.cell(inv_hdr_row, c_vec_col_idx).alignment = al_center

    inv_start_row = inv_hdr_row + 1
    inv_end_row = inv_start_row + n_free - 1
    col_k11_start = "C"
    col_k11_end = get_column_letter(3 + n_free - 1)

    for i_f in range(n_free):
        r = inv_start_row + i_f
        ws.row_dimensions[r].height = 20
        d_lbl_col_let = get_column_letter(d_lbl_col)
        ws.cell(r, 1, f"={d_lbl_col_let}{k11_data_start + i_f}")
        apply_styling(ws, r, r, 1, 1, fill=fill_white, font=f_bold, align=al_center)

        if r == inv_start_row + (n_free // 2):
            ws.cell(r, 2, "=")
            ws.cell(r, 2).alignment = al_center

        ws.cell(r, c_vec_col_idx, f"=A{k11_data_start + i_f}")
        apply_styling(ws, r, r, c_vec_col_idx, c_vec_col_idx, fill=fill_white, font=f_bold, align=al_right, num_fmt="#,##0.00")

    ws["C" + str(inv_start_row)] = ArrayFormula(
        ref=f"C{inv_start_row}:{col_k11_end}{inv_end_row}",
        text=f"=MINVERSE(C{k11_data_start}:{col_k11_end}{k11_end_row})"
    )
    apply_styling(ws, inv_start_row, inv_end_row, 3, 2 + n_free, fill=fill_white, font=f_regular, align=al_right, num_fmt="0.00000E+00")

    # 6. D=MMULT(MIN)*C
    du_banner_row = inv_end_row + 2
    ws.merge_cells(start_row=du_banner_row, start_column=3, end_row=du_banner_row, end_column=4)
    ws.cell(du_banner_row, 3, "D=MMULT(MIN)*C")
    apply_styling(ws, du_banner_row, du_banner_row, 3, 4, fill=fill_yellow_banner, font=f_yellow_banner, align=al_center)
    ws.row_dimensions[du_banner_row].height = 20

    du_hdr_row = du_banner_row + 1
    ws.cell(du_hdr_row, 1, "D")
    ws.cell(du_hdr_row, 1).font = f_bold
    ws.cell(du_hdr_row, 1).alignment = al_center

    ws.cell(du_hdr_row, 3, "MM*C")
    ws.cell(du_hdr_row, 3).font = f_bold
    ws.cell(du_hdr_row, 3).alignment = al_center

    du_start_row = du_hdr_row + 1
    du_end_row = du_start_row + n_free - 1

    for i_f in range(n_free):
        r = du_start_row + i_f
        ws.row_dimensions[r].height = 20
        ws.cell(r, 1, f"=A{inv_start_row + i_f}")
        apply_styling(ws, r, r, 1, 1, fill=fill_white, font=f_bold, align=al_center)

        if r == du_start_row + (n_free // 2):
            ws.cell(r, 2, "=")
            ws.cell(r, 2).alignment = al_center

        ws.cell(r, 4, "cm")
        apply_styling(ws, r, r, 4, 4, fill=fill_white, font=f_bold, align=al_center)

    ws["C" + str(du_start_row)] = ArrayFormula(
        ref=f"C{du_start_row}:C{du_end_row}",
        text=f"=MMULT(C{inv_start_row}:{col_k11_end}{inv_end_row}, {c_vec_col_let}{inv_start_row}:{c_vec_col_let}{inv_end_row})"
    )
    apply_styling(ws, du_start_row, du_end_row, 3, 3, fill=fill_white, font=f_bold, align=al_right, num_fmt="0.000000000")

    # Pedagogical Card 3
    card3_lines = [
        "• Submatriz [K11]: Extraída dinámicamente con INDEX/MATCH desde Karmadura usando los grados libres.",
        "• Vector de Cargas {C}: Fuerzas externas actuantes en los nudos libres.",
        "• Inversión [MINVERSA]: Matriz de flexibilidad [K11]⁻¹ en cm/Kgf.",
        "• Multiplicación [MMULT]: {D_libres} = [K11]⁻¹ · {C} obtiene los desplazamientos reales en cm."
    ]
    write_pedag_card(ws, k11_hdr_row, du_end_row, pedag_c_start, pedag_c_end,
                      "3. CÁLCULO DE DESPLAZAMIENTOS: {D} = [K₁₁]⁻¹ · {C}", card3_lines)

    # 7. ECUACION MATRICIAL REDUCIDA PARA EL CALCULO DE REACCIONES (K21)
    k21_hdr_row = du_end_row + 3
    ws.cell(k21_hdr_row, 1, "Ecuacion matricial reducida para el calculo de reacciones")
    ws.cell(k21_hdr_row, 1).font = f_prof_title_no_ul
    ws.row_dimensions[k21_hdr_row].height = 20

    k21_sub_hdr = k21_hdr_row + 1
    ws.cell(k21_sub_hdr, 1, "C")
    ws.cell(k21_sub_hdr, 1).font = f_bold
    ws.cell(k21_sub_hdr, 1).alignment = al_center

    du_col_dest = 3 + n_free + 1
    du_col_dest_let = get_column_letter(du_col_dest)
    ws.cell(k21_sub_hdr, du_col_dest, "D")
    ws.cell(k21_sub_hdr, du_col_dest).font = f_bold
    ws.cell(k21_sub_hdr, du_col_dest).alignment = al_center

    k21_data_start = k21_sub_hdr + 1
    k21_end_row = k21_data_start + n_rest - 1
    k21_col_end = get_column_letter(3 + n_free - 1)

    k21_helper_col = 3 + n_free + 2
    k21_helper_col_let = get_column_letter(k21_helper_col)

    for i_r in range(n_rest):
        r = k21_data_start + i_r
        ws.row_dimensions[r].height = 20
        k = i_r + 1

        ws.cell(r, k21_helper_col, f'=IFERROR(INDEX(C${disp_start_row}:C${disp_end_row}, MATCH({k}, G${disp_start_row}:G${disp_end_row}, 0)), "-")')
        apply_styling(ws, r, r, k21_helper_col, k21_helper_col, fill=fill_card, font=f_note, align=al_center)

        ws.cell(r, 1, f'=IF(${k21_helper_col_let}{r}="-", "-", IFERROR(INDEX(B${cargas_start_row}:B${cargas_end_row}, MATCH(${k21_helper_col_let}{r}, C${cargas_start_row}:C${cargas_end_row}, 0)), "-"))')
        apply_styling(ws, r, r, 1, 1, fill=fill_gray_reac, font=f_bold, align=al_center)

        if i_r == (n_rest // 2):
            ws.cell(r, 2, "=")
            ws.cell(r, 2).alignment = al_center

        for j_f in range(n_free):
            c_dest = 3 + j_f
            col_dest_let = get_column_letter(c_dest)
            formula = (
                f'=IF(OR(${k21_helper_col_let}{r}="-", {col_dest_let}${k11_col_dofs_row}="-"), 0, '
                f'IFERROR(INDEX($B${mat_karm_row}:${karm_end_col_let}${mat_karm_end}, MATCH(${k21_helper_col_let}{r}, ${lbl_col_let}${mat_karm_row}:${lbl_col_let}${mat_karm_end}, 0), MATCH({col_dest_let}${k11_col_dofs_row}, $B${hdr_karm_row}:${karm_end_col_let}${hdr_karm_row}, 0)), 0))'
            )
            ws.cell(r, c_dest, formula)
            apply_styling(ws, r, r, c_dest, c_dest, fill=fill_white, font=f_regular, align=al_right, num_fmt="#,##0.00")

        if i_r < n_free:
            ws.cell(r, du_col_dest, f"=C{du_start_row + i_r}")
            apply_styling(ws, r, r, du_col_dest, du_col_dest, fill=fill_white, font=f_regular, align=al_right, num_fmt="0.000000000")

    if n_rest < n_free:
        for extra_i in range(n_rest, n_free):
            r = k21_data_start + extra_i
            ws.row_dimensions[r].height = 20
            ws.cell(r, du_col_dest, f"=C{du_start_row + extra_i}")
            apply_styling(ws, r, r, du_col_dest, du_col_dest, fill=fill_white, font=f_regular, align=al_right, num_fmt="0.000000000")

    reac_res_start = max(k21_end_row, k21_data_start + n_free - 1) + 2
    reac_res_end = reac_res_start + n_rest - 1

    for i_r in range(n_rest):
        r = reac_res_start + i_r
        ws.row_dimensions[r].height = 20
        ws.cell(r, 1, f"=A{k21_data_start + i_r}")
        apply_styling(ws, r, r, 1, 1, fill=fill_white, font=f_bold, align=al_center)

        if i_r == (n_rest // 2):
            ws.cell(r, 2, "=")
            ws.cell(r, 2).alignment = al_center

        ws.cell(r, 4, "Kgf")
        apply_styling(ws, r, r, 4, 4, fill=fill_white, font=f_bold, align=al_center)

    ws["C" + str(reac_res_start)] = ArrayFormula(
        ref=f"C{reac_res_start}:C{reac_res_end}",
        text=f"=MMULT(C{k21_data_start}:{k21_col_end}{k21_end_row}, {du_col_dest_let}{k21_data_start}:{du_col_dest_let}{k21_data_start + n_free - 1})"
    )
    apply_styling(ws, reac_res_start, reac_res_end, 3, 3, fill=fill_white, font=f_bold, align=al_right, num_fmt="#,##0.00")

    # Pedagogical Card 4
    card4_lines = [
        "• Submatriz [K21]: Filas de apoyos (D=0) y columnas de grados libres (D≠0).",
        "• Multiplicación [MMULT]: {R} = [K21] · {D_libres} genera las reacciones en cada apoyo en Kgf.",
        "• Verificación de Equilibrio Global: ∑ Fx = 0  y  ∑ Fy = 0 (Equilibrio estático exacto)."
    ]
    write_pedag_card(ws, k21_hdr_row, reac_res_end, pedag_c_start, pedag_c_end,
                      "4. CÁLCULO DE REACCIONES: {R} = [K₂₁] · {D}", card4_lines)

    # 8. MATRIZ DE DESPLAZAMIENTOS TOTAL (12 DOFs)
    dtot_hdr_row = reac_res_end + 3
    ws.cell(dtot_hdr_row, 1, "Matriz de desplazamientos total")
    ws.cell(dtot_hdr_row, 1).font = f_prof_title_no_ul
    ws.row_dimensions[dtot_hdr_row].height = 20

    dtot_start_row = dtot_hdr_row + 2
    dtot_end_row = dtot_start_row + dofs_total - 1

    for idx_d, dof_name in enumerate(dofs_list):
        r = dtot_start_row + idx_d
        ws.row_dimensions[r].height = 19
        node_num = (idx_d // 2) + 1
        is_x = (idx_d % 2 == 0)

        ws.cell(r, 1, f"D{node_num}{'X' if is_x else 'Y'}")
        ws.cell(r, 2, f"U{node_num}{'X' if is_x else 'Y'}")
        apply_styling(ws, r, r, 1, 2, fill=fill_white, font=f_bold, align=al_center)

        ws.cell(r, 4, f'=IFERROR(INDEX(C${du_start_row}:C${du_end_row}, MATCH(A{r}, A${du_start_row}:A${du_end_row}, 0)), 0)')
        apply_styling(ws, r, r, 4, 4, fill=fill_white, font=f_bold, align=al_right, num_fmt="0.000000000")

    # 9. CALCULO DE LAS FUERZAS EN LAS 8 BARRAS
    forces_hdr_row = dtot_start_row + dofs_total + 2
    ws.cell(forces_hdr_row, 1, "Calculo de las fuerzas")
    ws.cell(forces_hdr_row, 1).font = f_prof_title_no_ul
    ws.row_dimensions[forces_hdr_row].height = 22

    f_bar_curr = forces_hdr_row + 2
    bar_force_result_rows = {}

    for m in range(1, MAX_BARS + 1):
        t1_row = 5 + m
        t2_row = t2_data_start + (m - 1)

        r_hdr1 = f_bar_curr
        ws.row_dimensions[r_hdr1].height = 18
        ws.cell(r_hdr1, 1, "AE/L")
        ws.cell(r_hdr1, 3, "cx")
        ws.cell(r_hdr1, 4, "cy")
        ws.cell(r_hdr1, 5, "cx")
        ws.cell(r_hdr1, 6, "cy")
        ws.cell(r_hdr1, 7, "D")
        apply_styling(ws, r_hdr1, r_hdr1, 1, 1, fill=fill_white, font=f_bold, align=al_center, border=None)
        apply_styling(ws, r_hdr1, r_hdr1, 3, 7, fill=fill_white, font=f_bold, align=al_center, border=None)

        r_hdr2 = f_bar_curr + 1
        ws.row_dimensions[r_hdr2].height = 16
        ws.cell(r_hdr2, 3, "-")
        ws.cell(r_hdr2, 4, "-")
        ws.cell(r_hdr2, 5, "+")
        ws.cell(r_hdr2, 6, "+")
        apply_styling(ws, r_hdr2, r_hdr2, 3, 6, fill=fill_white, font=f_bold, align=al_center, border=None)

        r3 = f_bar_curr + 2
        ws.row_dimensions[r3].height = 20
        ws.cell(r3, 1, f"=N{t1_row}")
        apply_styling(ws, r3, r3, 1, 1, fill=fill_white, font=f_bold, align=al_right, border=border_thin, num_fmt="#,##0.00")

        ws.cell(r3, 3, f'=IF(J{t1_row}="","", -J{t1_row})')
        ws.cell(r3, 4, f'=IF(K{t1_row}="","", -K{t1_row})')
        ws.cell(r3, 5, f'=IF(J{t1_row}="","", J{t1_row})')
        ws.cell(r3, 6, f'=IF(K{t1_row}="","", K{t1_row})')
        apply_styling(ws, r3, r3, 3, 6, fill=fill_white, font=f_regular, align=al_right, border=border_thin, num_fmt="0.0000")

        ws.cell(r3, 8, f'=IF(C{t2_row}="","", "D" & C{t2_row} & "X")')
        ws.cell(r3, 7, f'=IF(H{r3}="","", IFERROR(VLOOKUP(H{r3}, A${dtot_start_row}:D${dtot_end_row}, 4, FALSE), 0))')
        apply_styling(ws, r3, r3, 7, 7, fill=fill_white, font=f_regular, align=al_right, border=border_thin, num_fmt="0.00000000")
        apply_styling(ws, r3, r3, 8, 8, fill=fill_white, font=f_bold, align=al_center, border=border_thin)

        r4 = f_bar_curr + 3
        ws.row_dimensions[r4].height = 20
        ws.cell(r4, 1, f"F{m}=")
        ws.cell(r4, 1).font = f_bold
        ws.cell(r4, 1).alignment = al_center

        ws.cell(r4, 2, f'=IF(OR(A{r3}=0, C{t2_row}=""), 0, (MMULT(C{r3}:F{r3}, G{r3}:G{r3+3})*A{r3}))')
        apply_styling(ws, r4, r4, 2, 2, fill=fill_yellow_banner, font=f_bold, align=al_right, border=border_thin, num_fmt="#,##0.000")

        ws.cell(r4, 8, f'=IF(C{t2_row}="","", "D" & C{t2_row} & "Y")')
        ws.cell(r4, 7, f'=IF(H{r4}="","", IFERROR(VLOOKUP(H{r4}, A${dtot_start_row}:D${dtot_end_row}, 4, FALSE), 0))')
        apply_styling(ws, r4, r4, 7, 7, fill=fill_white, font=f_regular, align=al_right, border=border_thin, num_fmt="0.00000000")
        apply_styling(ws, r4, r4, 8, 8, fill=fill_white, font=f_bold, align=al_center, border=border_thin)

        r5 = f_bar_curr + 4
        ws.row_dimensions[r5].height = 20
        ws.cell(r5, 8, f'=IF(F{t2_row}="","", "D" & F{t2_row} & "X")')
        ws.cell(r5, 7, f'=IF(H{r5}="","", IFERROR(VLOOKUP(H{r5}, A${dtot_start_row}:D${dtot_end_row}, 4, FALSE), 0))')
        apply_styling(ws, r5, r5, 7, 7, fill=fill_white, font=f_regular, align=al_right, border=border_thin, num_fmt="0.00000000")
        apply_styling(ws, r5, r5, 8, 8, fill=fill_white, font=f_bold, align=al_center, border=border_thin)

        r6 = f_bar_curr + 5
        ws.row_dimensions[r6].height = 20
        ws.cell(r6, 8, f'=IF(F{t2_row}="","", "D" & F{t2_row} & "Y")')
        ws.cell(r6, 7, f'=IF(H{r6}="","", IFERROR(VLOOKUP(H{r6}, A${dtot_start_row}:D${dtot_end_row}, 4, FALSE), 0))')
        apply_styling(ws, r6, r6, 7, 7, fill=fill_white, font=f_regular, align=al_right, border=border_thin, num_fmt="0.00000000")
        apply_styling(ws, r6, r6, 8, 8, fill=fill_white, font=f_bold, align=al_center, border=border_thin)

        r7 = f_bar_curr + 6
        ws.row_dimensions[r7].height = 20
        ws.cell(r7, 2, "Kgf")
        ws.cell(r7, 2).font = f_bold
        ws.cell(r7, 2).alignment = al_center

        ws.merge_cells(start_row=r7, start_column=3, end_row=r7, end_column=6)
        ws.cell(r7, 3, f'=IF(C{t2_row}="","", IF(B{r4}>0.001,"BARRA A TRACCIÓN (+)",IF(B{r4}<-0.001,"BARRA A COMPRESIÓN (-)","BARRA DE FUERZA NULA")))')
        apply_styling(ws, r7, r7, 3, 6, fill=fill_prof_green, font=f_bold, align=al_center, border=border_thin)

        bar_force_result_rows[m] = (r4, r7)
        f_bar_curr = r7 + 2

    # Pedagogical Card 5
    card5_lines = [
        "• Fila de Signos [-cx, -cy, +cx, +cy]: Deformación axial relativa de la barra m.",
        "• Fuerza Axial F_m: Obtenida mediante =(MMULT(C:F, G:G)*A) vinculada a los nudos reales.",
        "• Diagnóstico de Estado: Tracción (+), Compresión (-) o Fuerza Nula."
    ]
    card5_end = min(forces_hdr_row + 25, f_bar_curr - 1)
    write_pedag_card(ws, forces_hdr_row, card5_end, pedag_c_start, pedag_c_end,
                      "5. CÁLCULO DE FUERZAS EN BARRAS: {Fₘ} = (AE/L) · [-cx, -cy, cx, cy] · {D}", card5_lines)

    # 10. RESUMEN GENERAL DE FUERZAS EN LAS BARRAS (8 BARRAS)
    sum_hdr_row = f_bar_curr + 1
    ws.merge_cells(start_row=sum_hdr_row, start_column=2, end_row=sum_hdr_row, end_column=8)
    ws.cell(sum_hdr_row, 2, "RESUMEN GENERAL DE FUERZAS EN LAS BARRAS (HASTA 8 BARRAS)")
    ws.cell(sum_hdr_row, 2).font = f_sec_hdr
    ws.cell(sum_hdr_row, 2).fill = fill_sec_teal
    ws.cell(sum_hdr_row, 2).alignment = al_left
    ws.row_dimensions[sum_hdr_row].height = 22

    sum_cols = [
        (sum_hdr_row + 1, sum_hdr_row + 1, 2, 2, "BARRA\n#"),
        (sum_hdr_row + 1, sum_hdr_row + 1, 3, 3, "NUDO\nINI"),
        (sum_hdr_row + 1, sum_hdr_row + 1, 4, 4, "NUDO\nFIN"),
        (sum_hdr_row + 1, sum_hdr_row + 1, 5, 5, "LONGITUD L\n(cm)"),
        (sum_hdr_row + 1, sum_hdr_row + 1, 6, 6, "ÁREA A\n(cm²)"),
        (sum_hdr_row + 1, sum_hdr_row + 1, 7, 7, "FUERZA AXIAL N\n(Kgf)"),
        (sum_hdr_row + 1, sum_hdr_row + 1, 8, 8, "ESTADO ESTRUCTURAL\n(TRACCIÓN / COMPRESIÓN)"),
    ]
    for r1, r2, c1, c2, txt in sum_cols:
        if r1 != r2 or c1 != c2:
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell = ws.cell(r1, c1)
        cell.value = txt
        cell.alignment = al_center

    apply_styling(ws, sum_hdr_row + 1, sum_hdr_row + 1, 2, 8, fill=fill_sec, font=f_tbl_hdr, align=al_center)
    ws.row_dimensions[sum_hdr_row + 1].height = 24

    for m in range(1, MAX_BARS + 1):
        r = sum_hdr_row + 1 + m
        ws.row_dimensions[r].height = 20
        t1_row = 5 + m
        t2_row = t2_data_start + (m - 1)
        r4_res, r7_res = bar_force_result_rows[m]

        bar_col_info = BAR_PASTEL_PALETTE[(m - 1) % len(BAR_PASTEL_PALETTE)]
        fill_bar_tracer = PatternFill("solid", fgColor=bar_col_info["bg"])

        ws.cell(r, 2, f'=IF(C{t2_row}="","", {m})')
        ws.cell(r, 3, f'=IF(C{t2_row}="","", C{t2_row})')
        ws.cell(r, 4, f'=IF(F{t2_row}="","", F{t2_row})')
        ws.cell(r, 5, f'=IF(I{t1_row}="","", I{t1_row})')
        ws.cell(r, 6, f'=IF(C{t2_row}="","", L{t1_row})')
        ws.cell(r, 7, f'=IF(C{t2_row}="","", B{r4_res})')
        ws.cell(r, 8, f'=IF(C{t2_row}="","", C{r7_res})')

        apply_styling(ws, r, r, 2, 2, fill=fill_bar_tracer, font=f_bold, align=al_center)
        apply_styling(ws, r, r, 3, 4, fill=fill_card, font=f_bold, align=al_center)
        apply_styling(ws, r, r, 5, 6, fill=fill_white, font=f_regular, align=al_right, num_fmt="#,##0.00")
        apply_styling(ws, r, r, 7, 7, fill=fill_yellow_banner, font=f_bold, align=al_right, num_fmt="#,##0.000")
        apply_styling(ws, r, r, 8, 8, fill=fill_prof_green, font=f_bold, align=al_center)

    print(f"Teacher Manual Sheet '{sheet_name}' (8 Bars, {MAX_NODES} Nodes, {dofs_total} DOFs) built successfully!")
